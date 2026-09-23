"""Run: .venv/bin/python -m unittest discover -s tests."""

import unittest
from io import BytesIO
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

from new_api_statistics.app import app
from new_api_statistics.report import (
    current_prices,
    price_details,
    decorate,
    export_excel,
    period,
    totals,
    rankings,
    TOKEN_FIELDS,
    convert_tokens,
    load_site_name,
    cost_formula,
    merge_failures,
)


class ReportTest(unittest.TestCase):
    def test_expression_prices_single_and_tiered(self):
        options = {
            "ModelRatio": {"claude-sonnet-5": 99},
            "billing_setting.billing_mode": {
                "claude-sonnet-5": "tiered_expr",
                "gpt-5.6-sol": "tiered_expr",
            },
            "billing_setting.billing_expr": {
                "claude-sonnet-5": 'tier("standard", p * 2 + cr * 0.2 + cc * 2.5 + cc1h * 4 + c * 10)',
                "gpt-5.6-sol": 'len <= 272000 ? tier("0_272k", p * 4 + cr * 0.4 + cc * 5 + c * 20) : tier("272k_plus", p * 8 + cr * 0.8 + cc * 10 + c * 30)',
            },
        }
        one = price_details("claude-sonnet-5", options)
        self.assertEqual(one["input_price"], Decimal(2))
        self.assertEqual(one["write_1h_price"], Decimal(4))
        self.assertEqual(one["output_price"], Decimal(10))
        self.assertEqual(one["pricing_mode"], "expression")
        self.assertEqual(one["price_tiers"][0]["name"], "standard")
        two = price_details("gpt-5.6-sol", options)
        self.assertIsNone(two["input_price"])
        self.assertEqual(
            [t["input_price"] for t in two["price_tiers"]], [Decimal(4), Decimal(8)]
        )
        self.assertEqual(two["price_tiers"][1]["condition"], "len > 272000")
        self.assertEqual(
            price_details(
                "gpt-6-astra",
                {
                    "billing_setting.billing_mode": {"gpt-6-astra": "tiered_expr"},
                    "billing_setting.billing_expr": {
                        "gpt-6-astra": 'tier("x", p * 1 + fixed(2))'
                    },
                },
            )["price_tiers"],
            [],
        )

    def test_expression_amount_and_usage_not_reconciled(self):
        original = dict(
            user_id=2,
            username="a",
            model_name="claude-sonnet-5",
            request_count=3,
            total_tokens=100,
            input_tokens=10,
            output_tokens=20,
            cache_read_tokens=40,
            cache_write_tokens=30,
            pricing_input_tokens=10,
            ratio_count=2,
            group_ratio=Decimal(4),
            amount=Decimal("1.25"),
        )
        options = {
            "billing_setting.billing_mode": {"claude-sonnet-5": "tiered_expr"},
            "billing_setting.billing_expr": {
                "claude-sonnet-5": 'tier("standard", p * 2 + cr * 0.2 + cc * 2.5 + cc1h * 4 + c * 10)'
            },
        }
        row = decorate([original], options)[0]
        self.assertEqual(row["cache_read_tokens"], 40)
        self.assertEqual(row["total_tokens"], 100)
        self.assertEqual(row["amount"], Decimal("1.25"))
        self.assertEqual(row["cost_formula"]["calculated"], Decimal("0.001212"))
        self.assertEqual(row["cost_formula"]["terms"][3]["tokens"], 30)
        self.assertEqual(row["cost_formula"]["terms"][3]["price"], Decimal("2.5"))
        wb = load_workbook(export_excel([row], "2026-09-01", "2026-09-01"))
        self.assertIn("表达式价格", wb.sheetnames)
        self.assertEqual(wb["表达式价格"]["D3"].value, 2)
        self.assertEqual(wb["表达式价格"]["G3"].value, 2.5)
        self.assertEqual(wb["表达式价格"]["H3"].value, "已拆分")
        self.assertEqual(wb["用户模型用量"]["P3"].value, 1.25)

    def test_single_tier_expression_matches_example(self):
        prices = price_details(
            "claude-sonnet-5",
            {
                "billing_setting.billing_mode": {"claude-sonnet-5": "tiered_expr"},
                "billing_setting.billing_expr": {
                    "claude-sonnet-5": 'tier("standard", p * 2 + cr * 0.2 + cc * 2.5 + cc1h * 4 + c * 10)'
                },
            },
        )
        row = dict(
            **prices,
            input_tokens=1623,
            output_tokens=226,
            cache_read_tokens=0,
            cache_write_tokens=0,
            group_ratio=Decimal(4),
            ratio_count=1,
            amount=Decimal("0.022024"),
        )
        formula = cost_formula(row)
        self.assertEqual(formula["calculated"], Decimal("0.022024"))
        self.assertEqual(formula["difference"], 0)
        row["cache_write_tokens"] = 100
        self.assertEqual(cost_formula(row)["calculated"], Decimal("0.023024"))

    def test_missing_tier_uses_low_price_for_trial(self):
        prices = price_details(
            "gpt-5.6-sol",
            {
                "billing_setting.billing_mode": {"gpt-5.6-sol": "tiered_expr"},
                "billing_setting.billing_expr": {
                    "gpt-5.6-sol": 'len <= 272000 ? tier("short", p * 4 + c * 20) : tier("long", p * 8 + c * 30)'
                },
            },
        )
        row = dict(
            **prices,
            input_tokens=100,
            output_tokens=10,
            cache_read_tokens=0,
            cache_write_tokens=0,
            group_ratio=Decimal(1),
            amount=Decimal("0.001"),
        )
        result = cost_formula(row)
        self.assertEqual(result["buckets"][0]["tier"], "short")
        self.assertTrue(result["buckets"][0]["inferred_low_tier"])
        self.assertEqual(result["calculated"], Decimal("0.0006"))

    def test_tier_rows_keep_actual_amounts_and_numeric_prices(self):
        options = {
            "billing_setting.billing_mode": {"gpt-5.6-sol": "tiered_expr"},
            "billing_setting.billing_expr": {
                "gpt-5.6-sol": 'len <= 272000 ? tier("short", p * 4 + cr * 0.4 + cc * 5 + c * 20) : tier("long", p * 8 + cr * 0.8 + cc * 10 + c * 30)'
            },
        }
        buckets = []
        for index, (tier, ratio, amount) in enumerate(
            [("short", "1", "0.001"), ("long", "2", "0.002"), ("", "1", "0.003")], 1
        ):
            buckets.append(
                dict(
                    matched_tier=tier,
                    group_ratio=ratio,
                    amount=amount,
                    request_count=1,
                    raw_input_tokens=105,
                    pricing_input_tokens=100,
                    input_tokens=100,
                    output_tokens=10,
                    cache_read_tokens=5,
                    cache_write_tokens=0,
                    total_tokens=115,
                    latest_at=index,
                    latest_id=index,
                )
            )
        source = dict(
            user_id=1,
            username="tester",
            display_name="",
            model_name="gpt-5.6-sol",
            token_id=0,
            token_name="",
            request_count=3,
            raw_input_tokens=315,
            pricing_input_tokens=300,
            input_tokens=300,
            output_tokens=30,
            cache_read_tokens=15,
            cache_write_tokens=0,
            total_tokens=345,
            ratio_count=2,
            group_ratio=Decimal(1),
            amount=Decimal("0.006"),
            tier_usage=buckets,
            failure_codes={"429": 1},
            failure_count=1,
        )
        rows = decorate([source], options)
        self.assertEqual([row["tier_name"] for row in rows], ["short", "long", "-"])
        self.assertEqual(
            [row["input_price"] for row in rows], [Decimal(4), Decimal(8), Decimal(4)]
        )
        self.assertEqual(
            [row["amount"] for row in rows],
            [Decimal("0.001"), Decimal("0.002"), Decimal("0.003")],
        )
        self.assertEqual(sum(row["request_count"] for row in rows), 3)
        self.assertEqual(totals(rows)["total_tokens"], 345)
        self.assertEqual(totals(rows)["amount"], Decimal("0.006"))
        self.assertEqual(sum(row["failure_count"] for row in rows), 1)
        self.assertEqual(rows[-1]["cost_formula"]["buckets"][0]["tier"], "short")
        self.assertTrue(rows[-1]["cost_formula"]["buckets"][0]["inferred_low_tier"])
        wb = load_workbook(export_excel(rows, "2026-09-22", "2026-09-22"))
        ws = wb["用户模型用量"]
        self.assertEqual(
            [ws[f"E{line}"].value for line in (3, 4, 5)], ["short", "long", "-"]
        )
        self.assertEqual([ws[f"L{line}"].value for line in (3, 4, 5)], [4, 8, 4])
        self.assertEqual(ws["P6"].value, "=SUM(P3:P5)")

    def test_unknown_recorded_tier_uses_current_low_price(self):
        options = {
            "billing_setting.billing_mode": {"gpt-6-astra": "tiered_expr"},
            "billing_setting.billing_expr": {
                "gpt-6-astra": 'len <= 272000 ? tier("0_272k", p * 10 + cr * 1 + cc * 12.5 + c * 50) : tier("272k_plus", p * 20 + cr * 2 + cc * 25 + c * 75)'
            },
        }
        source = dict(
            user_id=1, username="tester", display_name="", model_name="gpt-6-astra",
            token_id=0, token_name="", request_count=1, raw_input_tokens=4402,
            pricing_input_tokens=434, input_tokens=434, output_tokens=53,
            cache_read_tokens=3968, cache_write_tokens=0, total_tokens=4455,
            ratio_count=1, group_ratio=Decimal("3.4"), amount=Decimal("0.037258"),
            tier_usage=[dict(
                matched_tier="base", group_ratio="3.4", amount="0.037258",
                request_count=1, raw_input_tokens=4402, pricing_input_tokens=434,
                input_tokens=434, output_tokens=53, cache_read_tokens=3968,
                cache_write_tokens=0, total_tokens=4455, latest_at=1, latest_id=1,
            )],
            failure_codes={}, failure_count=0,
        )
        row = decorate([source], options)[0]
        self.assertEqual(row["tier_name"], "0_272k")
        self.assertEqual(row["input_price"], Decimal(10))
        self.assertEqual(row["output_price"], Decimal(50))
        self.assertEqual(row["cache_price"], Decimal(1))
        self.assertEqual(row["write_price"], Decimal("12.5"))
        bucket = row["cost_formula"]["buckets"][0]
        self.assertEqual(bucket["recorded_tier"], "base")
        self.assertEqual(bucket["tier"], "0_272k")
        self.assertTrue(bucket["inferred_low_tier"])
        self.assertEqual(bucket["calculated"], Decimal("0.0372572"))
        self.assertEqual(row["amount"], Decimal("0.037258"))

        # An obsolete name must not become its own row: it joins the current
        # low tier, while other current tiers and missing names remain separate.
        merged_source = deepcopy(source)
        original = source["tier_usage"][0]
        merged_source["tier_usage"] = [
            original,
            {**original, "matched_tier": "0_272k", "group_ratio": "3.4",
             "amount": "0.04", "latest_at": 2, "latest_id": 2},
            {**original, "matched_tier": "0_272k", "group_ratio": "2",
             "amount": "0.05", "latest_at": 3, "latest_id": 3},
            {**original, "matched_tier": "272k_plus", "group_ratio": "1",
             "amount": "0.08", "latest_at": 4, "latest_id": 4},
            {**original, "matched_tier": "", "group_ratio": "1",
             "amount": "0.01", "latest_at": 5, "latest_id": 5},
        ]
        merged_source["request_count"] = 5
        merged_source["amount"] = Decimal("0.217258")
        merged_source["failure_count"] = 1
        rows = decorate([merged_source], options)
        self.assertEqual([r["tier_name"] for r in rows], ["0_272k", "272k_plus", "-"])
        low = rows[0]
        self.assertEqual(low["request_count"], 3)
        self.assertEqual(low["total_tokens"], 13365)
        self.assertEqual(low["amount"], Decimal("0.127258"))
        self.assertEqual(low["group_ratio"], Decimal(2))
        self.assertEqual(low["ratio_count"], 2)
        self.assertEqual(len(low["cost_formula"]["buckets"]), 2)
        combined = low["cost_formula"]["buckets"][0]
        self.assertEqual(combined["request_count"], 2)
        self.assertEqual(combined["input_tokens"], 868)
        self.assertEqual(combined["actual"], Decimal("0.077258"))
        self.assertEqual(combined["calculated"], Decimal("0.0745144"))
        self.assertEqual(combined["terms"][0]["tokens"], 868)
        self.assertEqual(sum(r["request_count"] for r in rows), 5)
        self.assertEqual(totals(rows)["amount"], Decimal("0.217258"))
        self.assertEqual(sum(r["failure_count"] for r in rows), 1)
        ws = load_workbook(export_excel(rows, "2026-09-01", "2026-09-23"))["用户模型用量"]
        self.assertEqual([ws[f"E{n}"].value for n in (3, 4, 5)], ["0_272k", "272k_plus", "-"])
        self.assertEqual(ws["C3"].value, 3)
        self.assertEqual(ws["P3"].value, float(Decimal("0.127258")))

    def test_cost_formula_decimal_and_difference(self):
        row = dict(
            input_tokens=1000000,
            output_tokens=200000,
            cache_read_tokens=300000,
            cache_write_tokens=100000,
            input_price=Decimal(2),
            output_price=Decimal(10),
            cache_price=Decimal(".2"),
            write_price=Decimal("2.5"),
            group_ratio=Decimal(3),
            amount=Decimal("13"),
            ratio_count=2,
            converted_cache_read_tokens=None,
        )
        result = cost_formula(row)
        self.assertEqual(result["calculated"], Decimal("12.93"))
        self.assertEqual(result["difference"], Decimal(".07"))
        row["input_price"] = None
        self.assertIsNone(cost_formula(row)["calculated"])
        row["input_tokens"] = 0
        self.assertEqual(cost_formula(row)["calculated"], Decimal("6.93"))
        row["converted_cache_read_tokens"] = 0
        self.assertTrue(cost_formula(row)["converted"])

    def test_site_name_database_and_fallback(self):
        with patch("new_api_statistics.report.psycopg.connect") as connect:
            conn = connect.return_value.__enter__.return_value
            for value, expected in [
                (("示例网关",), "示例网关"),
                (None, "New API"),
                (("",), "New API"),
            ]:
                conn.execute.return_value.fetchone.return_value = value
                self.assertEqual(load_site_name(), expected)
            self.assertEqual(conn.execute.call_args.args[1], ("SystemName",))

    def test_site_name_rendered_and_escaped(self):
        with (
            patch("new_api_statistics.app.verify_admin", return_value=True),
            patch(
                "new_api_statistics.app.load_site_name", return_value="示例网关<script>"
            ),
        ):
            response = app.test_client().get("/statistics/", auth=("admin", "test"))
            text = response.get_data(as_text=True)
            self.assertIn("示例网关&lt;script&gt;", text)
            self.assertIn("用量统计</title>", text)
            self.assertIn("令牌汇总", text)
            self.assertIn("分令牌", text)
            self.assertIn(
                'data-model-mode="summary" aria-pressed="false">模型汇总', text
            )
            self.assertIn('data-model-mode="model" aria-pressed="true">分模型', text)
            self.assertIn("data-token-column hidden", text)
            self.assertIn('id="token-filter"', text)
            self.assertIn("筛选令牌", text)
            self.assertIn('id="group-filter"', text)
            self.assertIn("筛选分组", text)
            self.assertIn('data-column-toggle="tier_name" checked>档位', text)
            self.assertIn('<th data-column="tier_name">档位</th>', text)
            self.assertIn('data-preset="current-month">本月', text)
            self.assertIn('data-settings-tab="balance"', text)
            self.assertIn('data-settings-tab="notification"', text)
            self.assertIn('id="usage-channel-options"', text)
            script = Path(app.static_folder, "app.js").read_text()
            self.assertIn("filter-search", script)
            self.assertIn("row.style.display=matched?'':'none'", script)
            self.assertIn("输入关键字筛选", script)
            self.assertIn("function aggregateModels(data)", script)
            self.assertIn("modelMode==='summary'", script)
            self.assertIn("const modelSummaryColumns", script)
            self.assertIn("input.checked=false;input.disabled=true", script)
            presets = Path(app.static_folder, "presets.js").read_text()
            self.assertIn("preset === 'current-month'", presets)
            stylesheet = Path(app.static_folder, "app.css").read_text()
            self.assertIn(".filter-menu label[hidden]{display:none}", stylesheet)

    def setUp(self):
        self.rows = decorate(
            [
                dict(
                    user_id=1,
                    username="=danger",
                    display_name="测试显示名",
                    model_name="claude-a",
                    total_tokens=100,
                    request_count=2,
                    input_tokens=10,
                    output_tokens=20,
                    cache_write_tokens=30,
                    cache_read_tokens=40,
                    amount=Decimal("10"),
                    group_ratio=Decimal("3.8"),
                ),
                dict(
                    user_id=1,
                    username="=danger",
                    display_name="测试显示名",
                    model_name="gpt-b",
                    total_tokens=900,
                    request_count=4,
                    input_tokens=800,
                    output_tokens=100,
                    cache_write_tokens=0,
                    cache_read_tokens=500,
                    amount=Decimal("1"),
                    group_ratio=Decimal("1"),
                ),
            ],
            {},
        )

    def test_period_inclusive_shanghai(self):
        self.assertEqual(period("2026-07-26", "2026-08-25"), (1784995200, 1787673600))
        with self.assertRaises(ValueError):
            period("2026-08-25", "2026-07-26")
        with self.assertRaises(ValueError):
            period("bad", "2026-08-25")

    def test_token_breakdown(self):
        result = totals(self.rows)
        self.assertEqual([result[k] for k in TOKEN_FIELDS], [810, 120, 540, 30])
        self.assertEqual(result["total_tokens"], 1000)
        self.assertEqual(
            set(result),
            {
                *TOKEN_FIELDS,
                "total_tokens",
                "amount",
                "users",
                "models",
                "request_count",
            },
        )

    def test_failure_counts_merge_and_keep_failure_only_dimensions(self):
        rows = [
            dict(
                user_id=1,
                username="user",
                token_id=10,
                token_name="key",
                model_name="model-a",
                total_tokens=100,
            )
        ]
        failures = [
            dict(
                user_id=1,
                username="user",
                token_id=10,
                token_name="key-new",
                model_name="model-a",
                status_code="502",
                failure_count=1,
                latest_at=2,
            ),
            dict(
                user_id=1,
                username="user",
                token_id=10,
                token_name="key-new",
                model_name="model-a",
                status_code="429",
                failure_count=3,
                latest_at=3,
            ),
            dict(
                user_id=2,
                username="failure-only",
                token_id=20,
                token_name="failed-key",
                model_name="model-b",
                status_code="503",
                failure_count=2,
                latest_at=4,
            ),
        ]
        merged = merge_failures(rows, failures, by_token=True)
        by_model = {row["model_name"]: row for row in merged}
        self.assertEqual(by_model["model-a"]["failure_codes"], {"429": 3, "502": 1})
        self.assertEqual(by_model["model-a"]["failure_count"], 4)
        self.assertEqual(by_model["model-a"]["token_name"], "key-new")
        self.assertEqual(by_model["model-b"]["request_count"], 0)
        self.assertEqual(by_model["model-b"]["total_tokens"], 0)
        self.assertEqual(by_model["model-b"]["amount"], 0)

    def test_interval_average_rates(self):
        start, end = "2026-09-01T00:00:00", "2026-09-01T00:01:59"
        result = totals(self.rows, start, end)
        self.assertEqual(result["duration_seconds"], 120)
        self.assertEqual(result["tpm"], 500)
        self.assertEqual(result["rpm"], 3)
        single = totals(self.rows[:1], start, end)
        self.assertEqual(single["tpm"], 50)
        self.assertEqual(single["rpm"], 1)
        empty = totals([], start, end)
        self.assertEqual((empty["tpm"], empty["rpm"]), (0, 0))
        second = totals(self.rows, start, start)
        self.assertEqual((second["tpm"], second["rpm"]), (60000, 360))
        day = totals(self.rows, "2026-09-01", "2026-09-01")
        self.assertEqual(day["duration_seconds"], 86400)

    def test_rankings_aggregate_users_before_sorting(self):
        other = dict(
            self.rows[0],
            user_id=2,
            username="other",
            amount=Decimal("20"),
            total_tokens=50,
        )
        result = rankings([*self.rows, other])
        self.assertEqual(
            [r["username"] for r in result["user_tokens"]], ["=danger", "other"]
        )
        self.assertEqual(
            [r["username"] for r in result["user_amount"]], ["other", "=danger"]
        )
        self.assertEqual(result["user_tokens"][0]["total_tokens"], 1000)
        self.assertEqual(result["user_tokens"][0]["display_name"], "测试显示名")
        self.assertEqual(result["model_amount"][0]["amount"], 30)

    def test_time_boundaries(self):
        a, b = period("2026-07-26T18:20:30", "2026-07-26T18:20:30")
        self.assertEqual(b - a, 1)
        self.assertEqual(a, 1784995200 + 18 * 3600 + 20 * 60 + 30)
        self.assertEqual(period("2026-07-26 18:20", "2026-07-26T18:21:00")[1] - a, 31)
        with self.assertRaises(ValueError):
            period("2026-07-26T19:00:00", "2026-07-26T18:59:59")
        with self.assertRaises(ValueError):
            period("2026-07-26T19:00:00Z", "2026-07-27")

    def test_excel_exact_time(self):
        ws = load_workbook(
            export_excel([], "2026-07-26T18:20:30", "2026-07-26T19:21:31")
        ).active
        self.assertEqual(
            ws["B1"].value, "2026-07-26 18:20:30 至 2026-07-26 19:21:31（北京时间）"
        )

    def test_prices_missing_not_zero(self):
        options = {
            "ModelRatio": {"claude-a": 5},
            "CompletionRatio": {"claude-a": 5},
            "CacheRatio": {"claude-a": 0.1},
        }
        self.assertEqual(
            current_prices("claude-a", options),
            dict(input_price=10, output_price=50, cache_price=1, write_price=None),
        )

    def test_excel_numbers_merge_formulas_and_literal_names(self):
        ws = load_workbook(export_excel(self.rows, "2026-07-26", "2026-08-25")).active
        self.assertIn("A3:A4", str(ws.merged_cells))
        self.assertIn("B3:B4", str(ws.merged_cells))
        self.assertEqual(ws["A3"].data_type, "s")
        self.assertEqual(ws["A3"].value, "=danger")
        self.assertEqual(ws["B2"].value, "显示名")
        self.assertEqual(ws["B3"].value, "测试显示名")
        self.assertEqual(ws["C2"].value, "消费请求数")
        self.assertEqual(ws["K2"].value, "倍率")
        self.assertEqual(ws["E2"].value, "档位")
        self.assertEqual(ws["D3"].data_type, "s")
        self.assertEqual(ws["P3"].data_type, "n")
        self.assertEqual(ws["P5"].value, "=SUM(P3:P4)")
        for col in "CFGHIJP":
            self.assertEqual(ws[f"{col}3"].data_type, "n")
            self.assertEqual(ws[f"{col}5"].value, f"=SUM({col}3:{col}4)")
        self.assertEqual(ws.max_column, 16)
        self.assertNotIn("计价非缓存输入Token", [c.value for c in ws[2]])
        self.assertNotIn("折算状态", [c.value for c in ws[2]])

    def test_merged_conversion_used_consistently(self):
        original = dict(
            user_id=2,
            username="converted",
            model_name="gpt",
            request_count=2,
            total_tokens=10000,
            input_tokens=10000,
            output_tokens=0,
            cache_read_tokens=10000,
            cache_write_tokens=0,
            pricing_input_tokens=0,
            ratio_count=2,
            group_ratio=Decimal("3.4"),
            amount=Decimal(".032"),
        )
        options = {"ModelRatio": {"gpt": 1}, "CacheRatio": {"gpt": 0.5}}
        row = decorate([deepcopy(original)], options)[0]
        self.assertEqual(row["total_tokens"], row["converted_total_tokens"])
        self.assertEqual(row["cache_read_tokens"], row["converted_cache_read_tokens"])
        self.assertEqual(row["original_total_tokens"], 10000)
        self.assertEqual(row["original_cache_read_tokens"], 10000)
        self.assertIsInstance(row["cache_read_tokens"], int)
        self.assertIsInstance(row["total_tokens"], int)
        self.assertEqual(row["cache_read_tokens"], 9412)
        self.assertEqual(totals([row])["total_tokens"], row["total_tokens"])
        self.assertEqual(
            rankings([row])["user_tokens"][0]["total_tokens"], row["total_tokens"]
        )
        ws = load_workbook(export_excel([row], "2026-09-01", "2026-09-01")).active
        self.assertAlmostEqual(ws["F3"].value, float(row["total_tokens"]))
        self.assertAlmostEqual(ws["I3"].value, float(row["cache_read_tokens"]))
        self.assertEqual(ws["F3"].number_format, "#,##0")
        self.assertEqual(ws["I3"].number_format, "#,##0")
        self.assertEqual(ws["F2"].value, "总Token")
        self.assertAlmostEqual(ws["P3"].value, float(row["amount"]))
        self.assertNotIn("折算总Token", [c.value for c in ws[2]])
        for count in (1, 3):
            fallback = decorate([dict(original, ratio_count=count)], options)[0]
            self.assertEqual(fallback["total_tokens"], 10000)
        zero = decorate([dict(original, amount=Decimal(0))], options)[0]
        self.assertEqual(zero["total_tokens"], 0)

    def test_conversion_unique_solution_preserves_original(self):
        row = dict(
            ratio_count=2,
            group_ratio=Decimal("3.4"),
            amount=Decimal("0.032"),
            pricing_input_tokens=0,
            output_tokens=0,
            cache_write_tokens=0,
            cache_read_tokens=10000,
            total_tokens=10000,
            input_price=None,
            output_price=None,
            write_price=None,
            cache_price=Decimal(1),
        )
        convert_tokens(row)
        self.assertLess(row["converted_cache_read_tokens"], 10000)
        self.assertEqual(row["cache_read_tokens"], 10000)
        self.assertEqual(
            (
                row["converted_cache_read_tokens"]
                * row["cache_price"]
                * row["group_ratio"]
                / 1000000
            ).quantize(Decimal(".000001")),
            row["amount"],
        )
        row["group_ratio"] = Decimal(3)
        convert_tokens(row)
        self.assertGreater(row["converted_cache_read_tokens"], 10000)
        self.assertEqual(
            row["converted_total_tokens"], row["converted_cache_read_tokens"]
        )
        for count in (1, 3):
            row["ratio_count"] = count
            convert_tokens(row)
            self.assertIsNone(row["converted_total_tokens"])
        row.update(ratio_count=2, pricing_input_tokens=100000, input_price=Decimal(10))
        convert_tokens(row)
        self.assertIsNone(row["converted_total_tokens"])
        row["cache_price"] = None
        convert_tokens(row)
        self.assertIsNone(row["converted_total_tokens"])

    def test_excel_ranking_sheets(self):
        wb = load_workbook(export_excel(self.rows, "2026-07-26", "2026-08-25"))
        self.assertEqual(
            wb.sheetnames,
            ["用户模型用量", "模型消费", "用户Token用量", "用户消费", "区间汇总"],
        )
        self.assertEqual(wb["用户Token用量"]["C3"].value, "测试显示名")
        self.assertEqual(wb["用户Token用量"]["D3"].value, 1000)
        self.assertEqual(wb["用户消费"]["C3"].value, "测试显示名")
        self.assertEqual(wb["用户消费"]["D3"].value, 11)
        self.assertEqual(wb["用户消费"]["B3"].data_type, "s")
        self.assertEqual(wb["用户消费"]["C3"].data_type, "s")
        self.assertEqual(wb["用户消费"]["D4"].value, "=SUM(D3:D3)")
        summary = wb["区间汇总"]
        self.assertEqual(summary["B4"].value, "='用户模型用量'!F5")
        self.assertEqual(summary["B9"].value, "='用户模型用量'!C5")
        self.assertEqual(summary["B11"].value, "=B4/B10")
        self.assertEqual(summary["B12"].value, "=B9/B10")

    def test_empty_excel_no_circular_formula(self):
        ws = load_workbook(export_excel([], "2026-07-26", "2026-08-25")).active
        self.assertEqual(ws["F3"].value, 0)
        self.assertEqual(ws["P3"].value, 0)

    def test_developer_mode_does_not_change_excel(self):
        with (
            patch("new_api_statistics.app.verify_admin", return_value=True),
            patch("new_api_statistics.app.load_report", return_value=self.rows),
        ):
            books = []
            for suffix in ["", "&dev=0", "&dev=1", "&dev=2"]:
                response = app.test_client().get(
                    "/statistics/api/export?start=2026-07-26&end=2026-08-25" + suffix,
                    auth=("test_admin", "testing"),
                )
                self.assertEqual(response.status_code, 200)
                wb = load_workbook(BytesIO(response.data))
                self.assertEqual(wb.active.max_column, 16)
                books.append([[list(row) for row in ws.values] for ws in wb.worksheets])
            self.assertEqual(books[0], books[1])
            self.assertEqual(books[0], books[2])
            self.assertEqual(books[0], books[3])

    def test_developer_two_requests_failure_aggregation(self):
        with (
            patch("new_api_statistics.app.verify_admin", return_value=True),
            patch(
                "new_api_statistics.app.load_report", return_value=deepcopy(self.rows)
            ) as load,
        ):
            response = app.test_client().get(
                "/statistics/api/usage?start=2026-07-26&end=2026-08-25&dev=2",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.status_code, 200)
            load.assert_called_once_with(
                "2026-07-26", "2026-08-25", by_token=False, include_failures=True
            )

    def test_auth_and_user_model_filters(self):
        with patch(
            "new_api_statistics.app.verify_admin",
            side_effect=lambda u, p: u == "test_admin" and p == "testing",
        ):
            client = app.test_client()
            self.assertEqual(client.get("/statistics/api/usage").status_code, 401)
            with patch(
                "new_api_statistics.app.load_report", return_value=deepcopy(self.rows)
            ):
                response = client.get(
                    "/statistics/api/usage?start=2026-07-26&end=2026-08-25",
                    auth=("test_admin", "testing"),
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json["totals"]["amount"], "11")
                self.assertEqual(response.json["totals"]["request_count"], 6)
                response = client.get(
                    "/statistics/api/usage?user=absent&start=2026-07-26&end=2026-08-25",
                    auth=("test_admin", "testing"),
                )
                self.assertEqual(response.json["rows"], [])
                self.assertEqual(response.json["totals"]["rpm"], "0")
                response = client.get(
                    "/statistics/api/usage?model=gpt-b&start=2026-07-26&end=2026-08-25",
                    auth=("test_admin", "testing"),
                )
                self.assertEqual(len(response.json["rows"]), 1)
                self.assertEqual(response.json["rows"][0]["model_name"], "gpt-b")
                self.assertEqual(response.json["totals"]["amount"], "1")
                response = client.get(
                    "/statistics/api/export?user=%3Ddanger&model=claude-a&start=2026-07-26&end=2026-08-25",
                    auth=("test_admin", "testing"),
                )
                workbook = load_workbook(BytesIO(response.data))
                self.assertEqual(workbook.active["D3"].value, "claude-a")
                self.assertEqual(workbook.active.max_row, 4)
            self.assertEqual(
                client.get("/statistics/", auth=("admin", "testing")).status_code, 401
            )

    def test_token_detail_endpoint_is_separate_from_excel(self):
        token_rows = [dict(self.rows[0], token_id=7, token_name="key-a")]
        with (
            patch("new_api_statistics.app.verify_admin", return_value=True),
            patch(
                "new_api_statistics.app.load_report", return_value=token_rows
            ) as load,
        ):
            response = app.test_client().get(
                "/statistics/api/usage/by-token?start=2026-07-26&end=2026-08-25",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json["rows"][0]["token_name"], "key-a")
            load.assert_called_once_with("2026-07-26", "2026-08-25", by_token=True)

    def test_token_group_options_and_filtered_summary_endpoints(self):
        token_rows = [dict(token_id=7, token_name="key-a")]
        with (
            patch("new_api_statistics.app.verify_admin", return_value=True),
            patch(
                "new_api_statistics.app.load_token_options", return_value=token_rows
            ) as options,
            patch(
                "new_api_statistics.app.load_group_options",
                return_value=[{"group_name": "auto"}],
            ) as groups,
            patch("new_api_statistics.app.load_report", return_value=self.rows) as load,
        ):
            client = app.test_client()
            response = client.get(
                "/statistics/api/usage/tokens?start=2026-07-26&end=2026-08-25",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.json["rows"], token_rows)
            options.assert_called_once_with("2026-07-26", "2026-08-25")
            response = client.get(
                "/statistics/api/usage/groups?start=2026-07-26&end=2026-08-25",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.json["rows"], [{"group_name": "auto"}])
            groups.assert_called_once_with("2026-07-26", "2026-08-25")
            response = client.get(
                "/statistics/api/usage/by-selection?start=2026-07-26&end=2026-08-25&token_id=9&token_id=7&group=auto&by_token=1",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.status_code, 200)
            load.assert_called_once_with(
                "2026-07-26",
                "2026-08-25",
                by_token=True,
                token_ids=[7, 9],
                groups=["auto"],
            )
            response = client.get(
                "/statistics/api/usage/by-selection?start=2026-07-26&end=2026-08-25&token_id=bad",
                auth=("test_admin", "testing"),
            )
            self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()

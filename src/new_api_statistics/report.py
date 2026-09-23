"""Read-only database aggregates, current prices and numeric Excel output."""

import json
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import psycopg
from psycopg.rows import dict_row
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .expression_prices import extract_prices, low_tier, PRICE_VARS
from .historical_prices import (
    PRICE_FIELDS, matching_current_price, number, price_key, request_prices,
)

SQL = Path(__file__).with_name("usage.sql").read_text()
FAILURE_SQL = """WITH errors AS MATERIALIZED (
    SELECT id,created_at,user_id,username,token_id,token_name,model_name,
           COALESCE(NULLIF(COALESCE(NULLIF(btrim(other), ''), '{}')::jsonb
             ->>'status_code', ''), '未知') AS status_code
    FROM logs
    WHERE type=5 AND created_at >= %(start)s AND created_at < %(end)s
      /* scope_channels */
      AND (%(token_ids)s::bigint[] IS NULL OR token_id = ANY(%(token_ids)s::bigint[]))
      AND (%(groups)s::text[] IS NULL OR COALESCE("group", '') = ANY(%(groups)s::text[]))
), grouped AS (
    SELECT user_id,username,
           CASE WHEN %(by_token)s THEN token_id ELSE 0 END AS token_id,
           model_name,status_code,COUNT(*) AS failure_count,MAX(created_at) AS latest_at,
           (array_agg(COALESCE(NULLIF(btrim(token_name), ''), '未知令牌')
             ORDER BY created_at DESC,id DESC))[1] AS token_name
    FROM errors
    GROUP BY user_id,username,CASE WHEN %(by_token)s THEN token_id ELSE 0 END,
             model_name,status_code
)
SELECT * FROM grouped
ORDER BY username,user_id,token_id,model_name,status_code"""
TZ = ZoneInfo("Asia/Shanghai")
TOKEN_FIELDS = [
    "input_tokens",
    "output_tokens",
    "cache_read_tokens",
    "cache_write_tokens",
]
HEADERS = [
    "用户名",
    "显示名",
    "消费请求数",
    "模型",
    "档位",
    "总Token",
    "输入Token",
    "输出Token",
    "缓存读取Token",
    "缓存写入Token",
    "倍率",
    "输入价格（元/M）",
    "输出价格（元/M）",
    "缓存读价格（元/M）",
    "缓存写价格（元/M）",
    "消费金额（元）",
]
FIELDS = [
    "username",
    "display_name",
    "request_count",
    "model_name",
    "tier_name",
    "total_tokens",
    *TOKEN_FIELDS,
    "group_ratio",
    "input_price",
    "output_price",
    "cache_price",
    "write_price",
    "amount",
]


def convert_tokens(row):
    """Solve only for cache reads; never overwrite measured usage or money."""
    row.update(
        converted_cache_read_tokens=None,
        converted_total_tokens=None,
        conversion_status="不折算：非两种倍率",
    )
    if row.get("pricing_mode") == "expression":
        row["conversion_status"] = "不折算：表达式计费不能用单一价格反推用量"
        return
    if row.get("ratio_count") != 2:
        return
    ratio = row["group_ratio"]
    cache_price = row["cache_price"]
    if ratio is None or ratio <= 0 or cache_price is None or cache_price <= 0:
        row["conversion_status"] = "无法折算：倍率或缓存读价格无效"
        return
    fixed = Decimal(0)
    for token, price in [
        ("pricing_input_tokens", "input_price"),
        ("output_tokens", "output_price"),
        ("cache_write_tokens", "write_price"),
    ]:
        count = row.get(token)
        if count is None or count < 0 or (count and row[price] is None):
            row["conversion_status"] = "无法折算：单价缺失或计价用量无效"
            return
        fixed += Decimal(count) * (row[price] or Decimal(0))
    cache = (row["amount"] * 1000000 / ratio - fixed) / cache_price
    total = Decimal(row["total_tokens"]) - row["cache_read_tokens"] + cache
    if cache < 0 or total < 0:
        row["conversion_status"] = "无法折算：唯一解为负数"
        return
    row.update(
        converted_cache_read_tokens=cache,
        converted_total_tokens=total,
        conversion_status="已折算（非原始用量）",
    )


def convert_historical_row(row):
    """Reconcile cache per log-price/group bucket only when its ratio changed."""
    original_cache = row["cache_read_tokens"]
    converted = False
    statuses = []
    for bucket in row["pricing_buckets"]:
        old_cache = int(bucket["cache_read_tokens"])
        bucket["display_cache_read_tokens"] = old_cache
        bucket["display_total_tokens"] = int(bucket["total_tokens"])
        historical = bucket["historical_group_ratio"]
        current = bucket["current_group_ratio"]
        if historical is None or current is None:
            statuses.append("倍率缺失")
            continue
        if historical == current:
            continue
        if not bucket["current_match"]:
            statuses.append("历史价格不匹配当前档位")
            continue
        cache_price = row["cache_price"]
        if current <= 0 or cache_price is None or cache_price <= 0:
            statuses.append("当前倍率或缓存价格无效")
            continue
        fixed = Decimal(0)
        for count_field, price_field in (
            ("pricing_input_tokens", "input_price"),
            ("output_tokens", "output_price"),
            ("cache_write_tokens", "write_price"),
        ):
            count = int(bucket[count_field])
            price = row[price_field]
            if count < 0 or (count and price is None):
                statuses.append("其他计价单价缺失")
                break
            fixed += Decimal(count) * (price or Decimal(0))
        else:
            solved = (
                Decimal(str(bucket["amount"])) * 1000000 / current - fixed
            ) / cache_price
            if solved < 0 or not solved.is_finite():
                statuses.append("缓存读取唯一解无效")
                continue
            cache = int(solved.quantize(Decimal(1), rounding=ROUND_HALF_UP))
            bucket["display_cache_read_tokens"] = cache
            bucket["display_total_tokens"] += cache - old_cache
            converted = True
    cache = sum(bucket["display_cache_read_tokens"] for bucket in row["pricing_buckets"])
    row.update(
        converted_cache_read_tokens=cache if converted else None,
        converted_total_tokens=row["total_tokens"] - original_cache + cache if converted else None,
        conversion_status="；".join(statuses) if statuses else ("已折算" if converted else "倍率一致"),
    )


def parse_boundary(value, end=False):
    """Accept second/minute precision; preserve legacy whole-date queries."""
    try:
        if len(value) == 10:
            return datetime.combine(
                date.fromisoformat(value), time(23, 59, 59) if end else time(), TZ
            )
        for fmt in (
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%dT%H:%M",
            "%Y-%m-%d %H:%M",
        ):
            try:
                return datetime.strptime(value, fmt).replace(tzinfo=TZ)
            except ValueError:
                continue
    except (ValueError, TypeError):
        pass
    raise ValueError("请选择有效的开始、结束时间（北京时间）。")


def period(start, end):
    """Include the selected end second; SQL uses an exclusive upper bound."""
    first, last = parse_boundary(start), parse_boundary(end, end=True)
    if last < first:
        raise ValueError("结束时间不能早于开始时间。")
    if last - first >= timedelta(days=367):
        raise ValueError("单次查询最多 367 天，请缩小时间范围。")
    return int(first.timestamp()), int(last.timestamp()) + 1


def current_prices(model, options):
    """Use configured prices only: missing current prices remain unknown."""

    def value(key):
        raw = options.get(key, {}).get(model)
        return Decimal(str(raw)) if raw is not None else None

    ratio = value("ModelRatio")
    base = ratio * 2 if ratio is not None else None
    result = {"input_price": base}
    for field, key in [
        ("output_price", "CompletionRatio"),
        ("cache_price", "CacheRatio"),
        ("write_price", "CreateCacheRatio"),
    ]:
        multiplier = value(key)
        result[field] = (
            base * multiplier if base is not None and multiplier is not None else None
        )
    # Per-request pricing has no meaningful per-million-token price.
    if value("ModelPrice") is not None:
        result = dict.fromkeys(result)
    return result


def price_details(model, options):
    """Expression mode takes precedence over dormant legacy ratio settings."""
    if options.get("billing_setting.billing_mode", {}).get(model) != "tiered_expr":
        return {
            **current_prices(model, options),
            "pricing_mode": "ratio",
            "price_tiers": [],
        }
    tiers = extract_prices(options.get("billing_setting.billing_expr", {}).get(model))
    prices = dict.fromkeys([*list(PRICE_VARS.values())])
    if tiers and len(tiers) == 1:
        prices.update({field: tiers[0][field] for field in PRICE_VARS.values()})
    return {**prices, "pricing_mode": "expression", "price_tiers": tiers or []}


def expand_price_rows(rows, options):
    """Group expression usage by current tier; unknown names join the low tier."""
    expanded = []
    usage_fields = [
        "request_count",
        "raw_input_tokens",
        "pricing_input_tokens",
        "total_tokens",
        *TOKEN_FIELDS,
    ]
    for source in rows:
        details = price_details(source["model_name"], options)
        tiers = details["price_tiers"]
        usage = source.get("tier_usage") or []
        if usage and all("price_snapshot" in bucket for bucket in usage):
            expanded.extend(expand_historical_rows(source, details, options))
            continue
        if details["pricing_mode"] != "expression" or not tiers or not usage:
            row = dict(source)
            row.update(details)
            row["tier_name"] = "-"
            if details["pricing_mode"] != "expression":
                row.pop("tier_usage", None)
            if tiers:
                row.update(
                    {field: low_tier(tiers)[field] for field in PRICE_VARS.values()}
                )
            expanded.append(row)
            continue
        by_name = {tier["name"]: tier for tier in tiers}
        low_name = low_tier(tiers)["name"]
        grouped = {}
        for bucket in usage:
            recorded = bucket.get("matched_tier") or ""
            name = low_name if recorded and recorded not in by_name else recorded
            grouped.setdefault(name, []).append(bucket)
        order = {tier["name"]: index for index, tier in enumerate(tiers)}
        keys = sorted(
            grouped, key=lambda name: (order.get(name, len(tiers)), name == "", name)
        )
        failure_target = "" if "" in grouped else keys[0]
        for name in keys:
            buckets = grouped[name]
            row = dict(source)
            row.update(details)
            row["tier_name"] = name or "-"
            row["tier_usage"] = buckets
            for field in usage_fields:
                row[field] = sum(int(bucket[field]) for bucket in buckets)
            row["amount"] = sum(
                (Decimal(str(bucket["amount"])) for bucket in buckets), Decimal(0)
            )
            row["ratio_count"] = len({str(bucket["group_ratio"]) for bucket in buckets})
            latest = max(
                buckets,
                key=lambda bucket: (
                    bucket.get("latest_at", 0),
                    bucket.get("latest_id", 0),
                ),
            )
            row["group_ratio"] = Decimal(str(latest["group_ratio"]))
            tier = by_name.get(name) or low_tier(tiers)
            if tier:
                row.update({field: tier[field] for field in PRICE_VARS.values()})
            if name != failure_target:
                row["failure_codes"] = {}
                row["failure_count"] = 0
            expanded.append(row)
    return expanded


def expand_historical_rows(source, details, options):
    """Group requests by their observed unit prices, mapping matches to today."""
    candidates = details["price_tiers"] if details["pricing_mode"] == "expression" else []
    if details["pricing_mode"] == "ratio" and details["input_price"] is not None:
        candidates = [{"name": "-", **{field: details[field] for field in PRICE_FIELDS}}]
    current_groups = options.get("GroupRatio", {})
    grouped = {}
    for raw in source["tier_usage"]:
        bucket = dict(raw)
        historical = request_prices(bucket.get("price_snapshot"))
        match = matching_current_price(historical, candidates, bucket)
        if match:
            key = ("current", match["name"])
            display = {field: match[field] for field in PRICE_FIELDS}
            tier_name = match["name"]
        elif historical is not None:
            key = ("historical", price_key(historical))
            display = historical
            tier_name = "-"
        else:
            # An unpriced request must never borrow a tariff from another row.
            key = ("unknown", bucket.get("latest_id"))
            display = dict.fromkeys(PRICE_FIELDS)
            tier_name = "-"
        bucket["historical_prices"] = historical
        bucket["current_match"] = bool(match)
        bucket["current_group_ratio"] = number(current_groups.get(bucket.get("group_name")))
        bucket["historical_group_ratio"] = number(bucket.get("group_ratio"))
        if key not in grouped:
            grouped[key] = {"tier_name": tier_name, "prices": display, "buckets": []}
        grouped[key]["buckets"].append(bucket)

    usage_fields = ("request_count", "raw_input_tokens", "pricing_input_tokens",
                    "total_tokens", *TOKEN_FIELDS)
    ordered = sorted(
        grouped.values(),
        key=lambda group: (
            group["tier_name"] == "-",
            group["tier_name"],
            tuple("" if x is None else str(x) for x in price_key(group["prices"])),
        ),
    )
    result = []
    for index, item in enumerate(ordered):
        buckets = item["buckets"]
        row = dict(source)
        row.update(details)
        row.update(item["prices"])
        row["tier_name"] = item["tier_name"]
        row["pricing_buckets"] = buckets
        row["tier_usage"] = buckets
        for field in usage_fields:
            row[field] = sum(int(bucket[field]) for bucket in buckets)
        row["amount"] = sum(
            (Decimal(str(bucket["amount"])) for bucket in buckets), Decimal(0)
        )
        row["ratio_count"] = len({bucket["historical_group_ratio"] for bucket in buckets})
        latest = max(buckets, key=lambda b: (b.get("latest_at", 0), b.get("latest_id", 0)))
        row["group_ratio"] = latest["current_group_ratio"]
        if index:
            row["failure_codes"] = {}
            row["failure_count"] = 0
        result.append(row)
    return result


def cost_formula(row):
    """Explain displayed-token pricing without treating it as historical billing."""
    if row.get("pricing_mode") == "expression":
        tiers = row.get("price_tiers", [])
        source_buckets = row.get("tier_usage") or [
            {
                **{field: row[field] for field in TOKEN_FIELDS},
                "total_tokens": row.get(
                    "total_tokens", sum(row[field] for field in TOKEN_FIELDS)
                ),
                "request_count": row.get("request_count", 0),
                "matched_tier": "",
                "group_ratio": row.get("group_ratio"),
                "amount": row["amount"],
            }
        ]
        by_name = {tier["name"]: tier for tier in tiers}
        fallback = low_tier(tiers) if tiers else None
        buckets = []
        for source in source_buckets:
            matched = source.get("matched_tier") or ""
            tier = by_name.get(matched) or fallback
            inferred = tier is not None and matched not in by_name and len(tiers) > 1
            ratio = source.get("group_ratio")
            ratio = Decimal(str(ratio)) if ratio is not None else None
            actual = Decimal(str(source["amount"]))
            terms = []
            calculated = None
            if tier:
                # Omitted cache variables remain in p. Per the reporting rule,
                # all separately priced cache writes use the first (5m) rate.
                input_count = int(source["input_tokens"])
                read_count = int(source["cache_read_tokens"])
                write_count = int(source["cache_write_tokens"])
                if tier["cache_price"] is None:
                    input_count += read_count
                if tier["write_price"] is None:
                    input_count += write_count
                for label, count, price in [
                    ("计价输入", input_count, tier["input_price"]),
                    ("输出", int(source["output_tokens"]), tier["output_price"]),
                    (
                        "缓存读",
                        read_count if tier["cache_price"] is not None else 0,
                        tier["cache_price"],
                    ),
                    (
                        "缓存写",
                        write_count if tier["write_price"] is not None else 0,
                        tier["write_price"],
                    ),
                ]:
                    weighted = (
                        Decimal(count) * price
                        if price is not None
                        else (Decimal(0) if count == 0 else None)
                    )
                    terms.append(
                        dict(label=label, tokens=count, price=price, weighted=weighted)
                    )
                if ratio is not None and all(
                    term["weighted"] is not None for term in terms
                ):
                    calculated = (
                        sum((term["weighted"] for term in terms), Decimal(0))
                        * ratio
                        / 1000000
                    )
            buckets.append(
                {
                    "tier": tier["name"] if tier else matched or None,
                    "recorded_tier": matched or None,
                    "inferred_low_tier": inferred,
                    "request_count": int(source.get("request_count", 0)),
                    "total_tokens": int(source["total_tokens"]),
                    "input_tokens": int(source["input_tokens"]),
                    "output_tokens": int(source["output_tokens"]),
                    "cache_read_tokens": int(source["cache_read_tokens"]),
                    "cache_write_tokens": int(source["cache_write_tokens"]),
                    "ratio": ratio,
                    "terms": terms,
                    "calculated": calculated,
                    "actual": actual,
                }
            )
        # Logs can use several historical names for the same current tier.
        # Coalesce those buckets when the effective price and ratio are equal;
        # different ratios still require separate terms to remain auditable.
        merged = {}
        for bucket in buckets:
            key = (bucket["tier"], bucket["ratio"])
            if key not in merged:
                merged[key] = bucket
                continue
            target = merged[key]
            for field in (
                "request_count", "total_tokens", "input_tokens", "output_tokens",
                "cache_read_tokens", "cache_write_tokens", "actual",
            ):
                target[field] += bucket[field]
            target["inferred_low_tier"] |= bucket["inferred_low_tier"]
            if target["recorded_tier"] != bucket["recorded_tier"]:
                target["recorded_tier"] = None
            for term, extra in zip(target["terms"], bucket["terms"]):
                term["tokens"] += extra["tokens"]
                if term["weighted"] is None or extra["weighted"] is None:
                    term["weighted"] = None
                else:
                    term["weighted"] += extra["weighted"]
            if target["calculated"] is None or bucket["calculated"] is None:
                target["calculated"] = None
            else:
                target["calculated"] += bucket["calculated"]
        buckets = list(merged.values())
        calculated = (
            sum((bucket["calculated"] for bucket in buckets), Decimal(0))
            if buckets and all(bucket["calculated"] is not None for bucket in buckets)
            else None
        )
        return dict(
            mode="expression",
            tiers=tiers,
            buckets=buckets,
            terms=buckets[0]["terms"] if len(buckets) == 1 else [],
            ratio=row.get("group_ratio"),
            calculated=calculated,
            difference=row["amount"] - calculated if calculated is not None else None,
            ratio_count=row.get("ratio_count"),
        )
    terms = []
    ratio = row["group_ratio"]
    for label, token, price in [
        ("输入", "input_tokens", "input_price"),
        ("输出", "output_tokens", "output_price"),
        ("缓存读", "cache_read_tokens", "cache_price"),
        ("缓存写", "cache_write_tokens", "write_price"),
    ]:
        count, unit = row[token], row[price]
        weighted = (
            Decimal(count) * unit
            if unit is not None
            else (Decimal(0) if count == 0 else None)
        )
        terms.append(dict(label=label, tokens=count, price=unit, weighted=weighted))
    calculated = None
    if ratio is not None and all(t["weighted"] is not None for t in terms):
        calculated = sum((t["weighted"] for t in terms), Decimal(0)) * ratio / 1000000
    return dict(
        terms=terms,
        ratio=ratio,
        calculated=calculated,
        difference=row["amount"] - calculated if calculated is not None else None,
        converted=row.get("converted_cache_read_tokens") is not None,
        ratio_count=row.get("ratio_count"),
    )


def historical_cost_formula(row):
    """Explain current-tier reconciliation or unmatched historical-price billing."""
    prices = {field: row[field] for field in PRICE_FIELDS}
    grouped = {}
    for source in row["pricing_buckets"]:
        current = source["current_group_ratio"]
        historical = source["historical_group_ratio"]
        use_current = source["current_match"] and current is not None
        ratio = current if use_current else historical
        read = int(source["display_cache_read_tokens"])
        write = int(source["cache_write_tokens"])
        input_count = int(source["pricing_input_tokens"])
        if prices["cache_price"] is None:
            input_count += read
        if prices["write_price"] is None:
            input_count += write
        terms = []
        for label, count, price in (
            ("计价输入", input_count, prices["input_price"]),
            ("输出", int(source["output_tokens"]), prices["output_price"]),
            ("缓存读", read if prices["cache_price"] is not None else 0, prices["cache_price"]),
            ("缓存写", write if prices["write_price"] is not None else 0, prices["write_price"]),
        ):
            weighted = (Decimal(count) * price if price is not None else
                        Decimal(0) if count == 0 else None)
            terms.append(dict(label=label, tokens=count, price=price, weighted=weighted))
        calculated = (
            sum((term["weighted"] for term in terms), Decimal(0)) * ratio / 1000000
            if ratio is not None and all(term["weighted"] is not None for term in terms)
            else None
        )
        key = (ratio, use_current)
        if key not in grouped:
            grouped[key] = dict(
                ratio=ratio, current_ratio=use_current, terms=terms,
                request_count=int(source["request_count"]),
                calculated=calculated, actual=Decimal(str(source["amount"])),
            )
            continue
        target = grouped[key]
        target["request_count"] += int(source["request_count"])
        target["actual"] += Decimal(str(source["amount"]))
        for term, extra in zip(target["terms"], terms):
            term["tokens"] += extra["tokens"]
            term["weighted"] = (
                term["weighted"] + extra["weighted"]
                if term["weighted"] is not None and extra["weighted"] is not None
                else None
            )
        target["calculated"] = (
            target["calculated"] + calculated
            if target["calculated"] is not None and calculated is not None
            else None
        )
    buckets = list(grouped.values())
    calculated = (
        sum((bucket["calculated"] for bucket in buckets), Decimal(0))
        if buckets and all(bucket["calculated"] is not None for bucket in buckets)
        else None
    )
    return dict(
        mode="historical", prices=prices, buckets=buckets,
        calculated=calculated,
        difference=row["amount"] - calculated if calculated is not None else None,
        converted=row["converted_cache_read_tokens"] is not None,
        matched_tier=row["tier_name"] != "-",
    )


def decorate(rows, options):
    rows = expand_price_rows(rows, options)
    for row in rows:
        row["display_name"] = (row.get("display_name") or "").strip()
        for field in ["total_tokens", "request_count", *TOKEN_FIELDS]:
            row[field] = int(row[field])
        if "pricing_buckets" in row:
            convert_historical_row(row)
        else:
            convert_tokens(row)
        if "pricing_buckets" not in row and row["converted_cache_read_tokens"] is not None:
            cache = int(
                row["converted_cache_read_tokens"].quantize(
                    Decimal(1), rounding=ROUND_HALF_UP
                )
            )
            row["converted_cache_read_tokens"] = cache
            row["converted_total_tokens"] = (
                row["total_tokens"] - row["cache_read_tokens"] + cache
            )
        if row.get("pricing_input_tokens") is not None:
            row["pricing_input_tokens"] = int(row["pricing_input_tokens"])
        # Preserve source values for auditing; use one consistent report basis.
        for field in ("cache_read_tokens", "total_tokens"):
            row["original_" + field] = row[field]
            converted = row["converted_" + field]
            if converted is not None:
                row[field] = converted
        row["cost_formula"] = historical_cost_formula(row) if "pricing_buckets" in row else cost_formula(row)
    return rows


def load_site_name():
    """Read New API branding on page load so configuration changes take effect."""
    with psycopg.connect(
        connect_timeout=8,
        options="-c default_transaction_read_only=on -c statement_timeout=10000",
    ) as conn:
        row = conn.execute(
            "SELECT value FROM options WHERE key = %s", ("SystemName",)
        ).fetchone()
    return (row[0] or "").strip() if row and (row[0] or "").strip() else "New API"


def failure_key(row, by_token):
    key = (row["user_id"], row["username"], row["model_name"])
    return key + ((row.get("token_id"),) if by_token else ())


def merge_failures(rows, failures, by_token):
    """Attach status counts and retain dimensions that have only failed requests."""
    by_key = {failure_key(row, by_token): row for row in rows}
    latest = {}
    for failure in failures:
        key = failure_key(failure, by_token)
        row = by_key.get(key)
        if row is None:
            row = dict(
                user_id=failure["user_id"],
                username=failure["username"],
                token_id=failure["token_id"],
                token_name=failure["token_name"] if by_token else "",
                model_name=failure["model_name"],
                request_count=0,
                ratio_count=0,
                raw_input_tokens=0,
                input_tokens=0,
                output_tokens=0,
                cache_write_tokens=0,
                cache_read_tokens=0,
                pricing_input_tokens=0,
                total_tokens=0,
                amount=Decimal(0),
                group_ratio=None,
            )
            rows.append(row)
            by_key[key] = row
        row.setdefault("failure_codes", {})[failure["status_code"]] = int(
            failure["failure_count"]
        )
        previous_latest = latest.get(key)
        if by_token and (
            previous_latest is None or failure["latest_at"] > previous_latest
        ):
            row["token_name"] = failure["token_name"]
            latest[key] = failure["latest_at"]
    for row in rows:
        codes = row.setdefault("failure_codes", {})
        row["failure_codes"] = dict(
            sorted(
                codes.items(),
                key=lambda item: (
                    not item[0].isdigit(),
                    int(item[0]) if item[0].isdigit() else item[0],
                ),
            )
        )
        row["failure_count"] = sum(codes.values())
    rows.sort(
        key=lambda row: (
            row["username"],
            row["user_id"],
            row.get("token_name", ""),
            -row["total_tokens"],
            row["model_name"],
        )
    )
    return rows


def scoped_sql(query, channel_ids, excluded_channel_ids=None):
    """Restrict logs by channel IDs, never by the unrelated logs.group field.

    None means all channels; an empty array deliberately returns no records.
    The caller resolves channel IDs from the balance channel-tag inventory.
    """
    clause = (
        "AND (channel_id = ANY(%(channel_ids)s::bigint[]) OR "
        "(channel_id IS NULL AND array_position(%(channel_ids)s::bigint[], NULL) IS NOT NULL))"
        if channel_ids is not None
        else ""
    )
    if excluded_channel_ids is not None:
        clause += " AND (channel_id IS NULL OR NOT (channel_id = ANY(%(excluded_channel_ids)s::bigint[])))"
    return query.replace("/* scope_channels */", clause)


def load_report(
    start,
    end,
    *,
    by_token=False,
    token_ids=None,
    groups=None,
    include_failures=False,
    channel_ids=None,
    excluded_channel_ids=None,
):
    """Load either user-model totals or user-token-model totals read-only."""
    first, last = period(start, end)
    with psycopg.connect(
        connect_timeout=8,
        row_factory=dict_row,
        options="-c default_transaction_read_only=on -c statement_timeout=60000",
    ) as conn:
        conn.isolation_level = psycopg.IsolationLevel.REPEATABLE_READ
        rows = conn.execute(
            scoped_sql(SQL, channel_ids, excluded_channel_ids),
            {
                "start": first,
                "end": last,
                "by_token": by_token,
                "token_ids": token_ids,
                "groups": groups,
                "channel_ids": channel_ids,
                "excluded_channel_ids": excluded_channel_ids,
            },
        ).fetchall()
        if include_failures:
            failures = conn.execute(
                scoped_sql(FAILURE_SQL, channel_ids, excluded_channel_ids),
                {
                    "start": first,
                    "end": last,
                    "by_token": by_token,
                    "token_ids": token_ids,
                    "groups": groups,
                    "channel_ids": channel_ids,
                    "excluded_channel_ids": excluded_channel_ids,
                },
            ).fetchall()
            merge_failures(rows, failures, by_token)
        user_ids = list({row["user_id"] for row in rows})
        names = {}
        if user_ids:
            display_rows = conn.execute(
                "SELECT id, COALESCE(display_name, '') AS display_name FROM users WHERE id = ANY(%s)",
                (user_ids,),
            ).fetchall()
            names = {record["id"]: record["display_name"] for record in display_rows}
        for row in rows:
            row["display_name"] = names.get(row["user_id"], "")
        records = conn.execute(
            "SELECT key, value FROM options WHERE key = ANY(%s)",
            (
                [
                    "ModelRatio",
                    "CompletionRatio",
                    "CacheRatio",
                    "CreateCacheRatio",
                    "ModelPrice",
                    "GroupRatio",
                    "billing_setting.billing_mode",
                    "billing_setting.billing_expr",
                ],
            ),
        ).fetchall()
    options = {r["key"]: json.loads(r["value"] or "{}") for r in records}
    return decorate(rows, options)


def load_token_options(
    start, end, *, include_failures=False, channel_ids=None, excluded_channel_ids=None
):
    """Return the latest name of each token used in the selected interval."""
    first, last = period(start, end)
    with psycopg.connect(
        connect_timeout=8,
        row_factory=dict_row,
        options="-c default_transaction_read_only=on -c statement_timeout=60000",
    ) as conn:
        return conn.execute(
            """SELECT DISTINCT ON (token_id) token_id,
                      COALESCE(NULLIF(btrim(token_name), ''), '未知令牌') AS token_name
               FROM logs
               WHERE type = ANY(%s) AND created_at >= %s AND created_at < %s
                 AND (%s::bigint[] IS NULL OR channel_id = ANY(%s::bigint[])
                      OR (channel_id IS NULL AND array_position(%s::bigint[], NULL) IS NOT NULL))
                 AND (%s::bigint[] IS NULL OR channel_id IS NULL OR NOT (channel_id = ANY(%s::bigint[])))
               ORDER BY token_id, created_at DESC, id DESC""",
            (
                [2, 5] if include_failures else [2],
                first,
                last,
                channel_ids,
                channel_ids,
                channel_ids,
                excluded_channel_ids,
                excluded_channel_ids,
            ),
        ).fetchall()


def load_group_options(
    start, end, *, include_failures=False, channel_ids=None, excluded_channel_ids=None
):
    """Return groups used in the selected interval."""
    first, last = period(start, end)
    with psycopg.connect(
        connect_timeout=8,
        row_factory=dict_row,
        options="-c default_transaction_read_only=on -c statement_timeout=60000",
    ) as conn:
        return conn.execute(
            """SELECT DISTINCT COALESCE("group", '') AS group_name
               FROM logs
               WHERE type = ANY(%s) AND created_at >= %s AND created_at < %s
                 AND (%s::bigint[] IS NULL OR channel_id = ANY(%s::bigint[])
                      OR (channel_id IS NULL AND array_position(%s::bigint[], NULL) IS NOT NULL))
                 AND (%s::bigint[] IS NULL OR channel_id IS NULL OR NOT (channel_id = ANY(%s::bigint[])))
               ORDER BY group_name""",
            (
                [2, 5] if include_failures else [2],
                first,
                last,
                channel_ids,
                channel_ids,
                channel_ids,
                excluded_channel_ids,
                excluded_channel_ids,
            ),
        ).fetchall()


def totals(rows, start=None, end=None):
    tokens = sum(r["total_tokens"] for r in rows)
    amount = sum((r["amount"] for r in rows), Decimal(0))
    result = {
        "total_tokens": tokens,
        "amount": amount,
        "request_count": sum(r["request_count"] for r in rows),
        "users": len({r["user_id"] for r in rows}),
        "models": len({r["model_name"] for r in rows}),
        **{field: sum(r[field] for r in rows) for field in TOKEN_FIELDS},
    }
    if start is not None and end is not None:
        first, last = period(start, end)
        # Use the entire selected interval, including idle time and the end second.
        seconds = last - first
        result.update(
            duration_seconds=seconds,
            tpm=Decimal(tokens) * 60 / seconds,
            rpm=Decimal(result["request_count"]) * 60 / seconds,
        )
    return result


def rankings(rows):
    """Aggregate before ranking; normalized totals avoid counting caches twice."""
    users, models = {}, {}
    for row in rows:
        key = (row["user_id"], row["username"])
        user = users.setdefault(
            key,
            dict(
                user_id=key[0],
                username=key[1],
                display_name=row.get("display_name", ""),
                amount=Decimal(0),
                total_tokens=0,
                **dict.fromkeys(TOKEN_FIELDS, 0),
            ),
        )
        for field in ["amount", "total_tokens", *TOKEN_FIELDS]:
            user[field] += row[field]
        models[row["model_name"]] = (
            models.get(row["model_name"], Decimal(0)) + row["amount"]
        )
    return {
        "model_amount": sorted(
            (dict(model_name=k, amount=v) for k, v in models.items()),
            key=lambda r: (-r["amount"], r["model_name"]),
        ),
        "user_tokens": sorted(
            users.values(),
            key=lambda r: (-r["total_tokens"], r["username"], r["user_id"]),
        ),
        "user_amount": sorted(
            users.values(), key=lambda r: (-r["amount"], r["username"], r["user_id"])
        ),
    }


def export_excel(rows, start, end):
    """Merge user cells, retain numeric prices and calculate totals with SUM."""
    wb = Workbook()
    ws = wb.active
    ws.title = "用户模型用量"
    first, last = parse_boundary(start), parse_boundary(end, end=True)
    ws.append(
        [
            "统计时间",
            f"{first:%Y-%m-%d %H:%M:%S} 至 {last:%Y-%m-%d %H:%M:%S}（北京时间）",
        ]
    )
    ws.merge_cells("C1:P1")
    ws.append(HEADERS)
    for r in rows:
        values = [r.get(f) for f in FIELDS]
        ws.append(values)
        # Treat database names literally, including strings starting with '='.
        for col in (1, 2, 4, 5):
            ws.cell(ws.max_row, col).data_type = "s"
    last = ws.max_row
    ws.append(["总计"])
    for col in (3, 6, 7, 8, 9, 10, 16):
        letter = get_column_letter(col)
        ws.cell(last + 1, col, f"=SUM({letter}3:{letter}{last})" if rows else 0)
    begin = 3
    for i in range(1, len(rows) + 1):
        if i == len(rows) or (rows[i]["user_id"], rows[i]["username"]) != (
            rows[i - 1]["user_id"],
            rows[i - 1]["username"],
        ):
            if i + 2 > begin:
                ws.merge_cells(
                    start_row=begin, end_row=i + 2, start_column=1, end_column=1
                )
                ws.merge_cells(
                    start_row=begin, end_row=i + 2, start_column=2, end_column=2
                )
            begin = i + 3
    for cells in ws.iter_rows(min_row=3):
        for c in cells:
            c.alignment = Alignment(vertical="center")
            if c.column not in (1, 2, 4, 5):
                c.number_format = (
                    "#,##0" if c.column in (3, 6, 7, 8, 9, 10) else "#,##0.######"
                )
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="156B59")
    for i, width in enumerate(
        [24, 24, 18, 36, 18, 30, 22, 22, 30, 22, 12, 24, 24, 24, 24, 22], 1
    ):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:P{last}"
    ranked = rankings(rows)
    for title, key, columns in [
        (
            "模型消费",
            "model_amount",
            [("模型", "model_name"), ("消费金额（元）", "amount")],
        ),
        (
            "用户Token用量",
            "user_tokens",
            [
                ("用户名", "username"),
                ("显示名", "display_name"),
                ("总Token", "total_tokens"),
                ("输入Token", "input_tokens"),
                ("输出Token", "output_tokens"),
                ("缓存读取Token", "cache_read_tokens"),
                ("缓存写入Token", "cache_write_tokens"),
            ],
        ),
        (
            "用户消费",
            "user_amount",
            [
                ("用户名", "username"),
                ("显示名", "display_name"),
                ("消费金额（元）", "amount"),
            ],
        ),
    ]:
        sheet = wb.create_sheet(title)
        sheet.append(["统计时间", ws["B1"].value])
        if key == "user_tokens":
            sheet.cell(1, 3, ws["C1"].value)
            sheet.merge_cells("C1:H1")
        sheet.append(["序号", *[label for label, _ in columns]])
        for rank, row in enumerate(ranked[key], 1):
            sheet.append([rank, *[row[field] for _, field in columns]])
            for col, (_, field) in enumerate(columns, 2):
                if field in ("username", "display_name"):
                    sheet.cell(sheet.max_row, col).data_type = "s"
        end_row = sheet.max_row
        sheet.append(["总计"])
        for col, (_, field) in enumerate(columns, 2):
            if field not in ("username", "display_name"):
                letter = get_column_letter(col)
                sheet.cell(
                    end_row + 1,
                    col,
                    f"=SUM({letter}3:{letter}{end_row})" if ranked[key] else 0,
                )
        for col, (_, field) in enumerate(columns, 2):
            sheet.column_dimensions[get_column_letter(col)].width = (
                36 if field == "username" else 24
            )
            for line in range(3, sheet.max_row + 1):
                sheet.cell(line, col).number_format = (
                    "#,##0.######" if field == "amount" else "#,##0"
                )
        for c in sheet[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="156B59")
        sheet.freeze_panes = "D3" if key.startswith("user_") else "C3"
        sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns) + 1)}{end_row}"
    summary = wb.create_sheet("区间汇总")
    summary.append(["统计时间", ws["B1"].value])
    summary.cell(1, 3, ws["C1"].value)
    summary.column_dimensions["C"].width = 60
    summary.append(["指标", "数值"])
    for title, col in [
        ("消费金额（元）", "P"),
        ("总Token", "F"),
        ("输入Token", "G"),
        ("输出Token", "H"),
        ("缓存读取Token", "I"),
        ("缓存写入Token", "J"),
        ("消费请求数", "C"),
    ]:
        summary.append([title, f"='用户模型用量'!{col}{last + 1}"])
    first_second, last_second = period(start, end)
    summary.append(["区间分钟数", f"=({last_second}-{first_second})/60"])
    summary.append(["区间平均 TPM", "=B4/B10"])
    summary.append(["区间平均 RPM", "=B9/B10"])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 62
    for row in range(3, 13):
        summary.cell(row, 2).number_format = (
            "#,##0" if row in (4, 5, 6, 7, 8, 9) else "#,##0.######"
        )
    for c in summary[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="156B59")
    summary.freeze_panes = "B3"
    expression_rows = {
        row["model_name"]: row
        for row in rows
        if row.get("pricing_mode") == "expression"
    }
    if expression_rows:
        prices = wb.create_sheet("表达式价格")
        prices.append(
            ["当前表达式配置的单价；元/百万 Token。历史消费金额仍以日志为准。"]
        )
        prices.append(
            [
                "模型",
                "档位",
                "适用条件",
                "输入价格",
                "输出价格",
                "缓存读价格",
                "缓存写价格（5分钟）",
                "拆分状态",
            ]
        )
        for model, row in sorted(expression_rows.items()):
            for tier in row["price_tiers"] or [None]:
                prices.append(
                    [
                        model,
                        tier["name"] if tier else None,
                        tier["condition"] if tier else None,
                        *(
                            [
                                tier[field]
                                for field in (
                                    "input_price",
                                    "output_price",
                                    "cache_price",
                                    "write_price",
                                )
                            ]
                            if tier
                            else [None] * 4
                        ),
                        "已拆分" if tier else "表达式无法安全拆分",
                    ]
                )
                prices.cell(prices.max_row, 1).data_type = "s"
                for column in range(4, 8):
                    prices.cell(prices.max_row, column).number_format = "#,##0.######"
        for cell_ in prices[2]:
            cell_.font = Font(bold=True, color="FFFFFF")
            cell_.fill = PatternFill("solid", fgColor="156B59")
        for column in range(1, 9):
            prices.column_dimensions[get_column_letter(column)].width = 26
        prices.freeze_panes = "D3"
        prices.auto_filter.ref = f"A2:H{prices.max_row}"
    result = BytesIO()
    wb.save(result)
    result.seek(0)
    return result
